from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from accounting_agent.matching import InvoiceRef
from accounting_agent.payments import TRANSFER_PAID, UNPAID, InvoiceStatus
from accounting_agent.report import write_report

INVOICE = InvoiceRef(
    "KSEF-1",
    "509417000926",
    date(2026, 9, 9),
    "T-Mobile Polska S.A.",
    "5261040567",
    Decimal("246.28"),
    "PLN",
    "Vat",
)
DEBIT = {
    "booking_date": "2026-09-11",
    "operation_date": "2026-09-11",
    "amount": 246.28,
    "currency": "PLN",
    "kind": "przelew",
    "counterparty": "T-mobile Polska S.A.",
    "account": "",
    "description": "Tytuł: 509417000926",
    "method": "numer faktury",
    "note": "",
    "invoice": {
        "ksef_number": "KSEF-1",
        "number": "509417000926",
        "issue_date": "2026-09-09",
        "seller": "T-Mobile Polska S.A.",
        "gross": 246.28,
    },
}
RECORD = {
    "report": "x.xlsx",
    "statements": 1,
    "unreadable": [],
    "invoice_count": 1,
    "accounts": [
        {
            "name": "Alior",
            "bank": "Alior",
            "account": "58249000050000453016739104",
            "currency": "PLN",
            "period_from": "2026-09-01",
            "period_to": "2026-09-30",
            "files": 30,
        }
    ],
    "banks": [{"name": "Alior", "matched": [DEBIT], "missing": [], "other": []}],
}


def test_report_has_bank_and_invoice_sheets(tmp_path):
    path = tmp_path / "raport.xlsx"
    statuses = [
        InvoiceStatus(
            INVOICE, TRANSFER_PAID, "Alior, 11-10-2026", "przelew", "2026-10-11", "Alior", "2026-10"
        ),
    ]
    write_report(path, "Porównanie – Wrzesień 2026", RECORD, statuses)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Podsumowanie", "Alior", "Faktury – płatności"]
    invoices = wb["Faktury – płatności"]
    assert invoices["G5"].value == "zapłacona przelewem (w kolejnym miesiącu)"
    assert invoices["J5"].value == "2026-10"
    assert wb["Alior"]["G6"].value == "509417000926"


def test_report_lists_unpaid_invoice(tmp_path):
    path = tmp_path / "raport.xlsx"
    write_report(
        path,
        "t",
        {**RECORD, "banks": [{**RECORD["banks"][0], "matched": []}]},
        [InvoiceStatus(INVOICE, UNPAID, "wg faktury: przelew, termin 23-09-2026")],
    )
    invoices = load_workbook(path)["Faktury – płatności"]
    assert invoices["G5"].value == UNPAID
    assert invoices["K5"].value == "wg faktury: przelew, termin 23-09-2026"
