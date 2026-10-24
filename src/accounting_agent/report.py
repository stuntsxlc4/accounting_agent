"""Zestawienie porównania faktur z wyciągami w Excelu.

Strona wyciągów pochodzi z zapisanego wyniku porównania (record.py), a strona faktur
z rejestru płatności, więc zestawienie można przeliczyć bez ponownego czytania wyciągów.
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounting_agent.formatting import bank_account
from accounting_agent.payments import (
    CARD_PAID,
    CASH,
    CONFLICT,
    CORRECTED_TO_ZERO,
    CORRECTION,
    TO_EXPLAIN,
    TRANSFER_PAID,
    UNPAID,
    InvoiceStatus,
)

MONEY = "#,##0.00"
DATE = "DD-MM-YYYY"
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
SECTION = Font(bold=True, size=12)
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
PAID_FILL = PatternFill("solid", fgColor="E2EFDA")
MISSING_FILL = PatternFill("solid", fgColor="F8CBAD")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
NEUTRAL_FILL = PatternFill("solid", fgColor="F2F2F2")

# Kolejność i kolory statusów w arkuszu faktur
STATUS_FILLS = {
    TO_EXPLAIN: MISSING_FILL,
    CONFLICT: MISSING_FILL,
    UNPAID: WARNING_FILL,
    TRANSFER_PAID: PAID_FILL,
    CARD_PAID: PAID_FILL,
    CASH: PAID_FILL,
    CORRECTED_TO_ZERO: NEUTRAL_FILL,
    CORRECTION: NEUTRAL_FILL,
}

DEBIT_COLUMNS = [
    ("Data obciążenia", 14),
    ("Data operacji", 14),
    ("Kwota obciążenia", 16),
    ("Waluta", 8),
    ("Rodzaj", 16),
    ("Kontrahent / opis operacji", 60),
    ("Faktura – numer", 24),
    ("Faktura – data", 14),
    ("Faktura – sprzedawca", 32),
    ("Faktura – brutto", 16),
    ("Dopasowanie", 22),
    ("Uwagi", 40),
]

INVOICE_COLUMNS = [
    ("Data faktury", 14),
    ("Numer", 30),
    ("Sprzedawca", 42),
    ("NIP", 14),
    ("Brutto", 14),
    ("Waluta", 8),
    ("Status płatności", 26),
    ("Data zapłaty", 14),
    ("Bank", 14),
    ("Z wyciągu za", 14),
    ("Uwagi", 46),
]


def write_report(path: Path, title: str, record: dict, statuses: list[InvoiceStatus]) -> None:
    wb = Workbook()
    _summary_sheet(wb.active, title, record, statuses)
    for bank in record["banks"]:
        account = next(a for a in record["accounts"] if a["name"] == bank["name"])
        _bank_sheet(wb.create_sheet(bank["name"]), bank, account)
    _invoices_sheet(wb.create_sheet("Faktury – płatności"), statuses)
    wb.save(path)


def _summary_sheet(ws, title: str, record: dict, statuses: list[InvoiceStatus]) -> None:
    ws.title = "Podsumowanie"
    ws["A1"] = title
    ws["A1"].font = TITLE

    ws["A3"] = "Wyciągi bankowe"
    ws["A3"].font = SECTION
    _header_row(
        ws, 4, ["Rachunek", "Obciążeń", "Z fakturą", "BRAK FAKTURY", "Bez faktury z natury"]
    )
    row = 5
    for bank in record["banks"]:
        values = [
            bank["name"],
            len(bank["matched"]) + len(bank["missing"]) + len(bank["other"]),
            len(bank["matched"]),
            len(bank["missing"]),
            len(bank["other"]),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
        if bank["missing"]:
            ws.cell(row=row, column=4).fill = MISSING_FILL
        row += 1
    for message in record["unreadable"]:
        ws.cell(row=row, column=1, value="Nie odczytano").font = BOLD
        ws.cell(row=row, column=2, value=message).fill = MISSING_FILL
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Faktury miesiąca").font = SECTION
    row += 1
    _header_row(ws, row, ["Status płatności", "Faktur"])
    for status, fill in STATUS_FILLS.items():
        count = sum(1 for s in statuses if s.status == status)
        if count:
            row += 1
            ws.cell(row=row, column=1, value=status).fill = fill
            ws.cell(row=row, column=2, value=count)
    row += 2
    ws.cell(row=row, column=1, value="Stan na").font = BOLD
    ws.cell(row=row, column=2, value=date.today()).number_format = DATE
    _widths(ws, [30, 14, 14, 16, 22])


def _bank_sheet(ws, bank: dict, account: dict) -> None:
    ws["A1"] = (
        f"{bank['name']} – rachunek {bank_account(account['account'])} ({account['currency']})"
    )
    ws["A1"].font = TITLE
    if account["period_from"]:
        ws["A2"] = (
            f"Okres: {_date(account['period_from']):%d-%m-%Y} – "
            f"{_date(account['period_to']):%d-%m-%Y}, plików: {account['files']}"
        )

    sections = [
        ("Obciążenia z fakturą", bank["matched"], PAID_FILL, True),
        ("Obciążenia BEZ FAKTURY – do wyjaśnienia", bank["missing"], MISSING_FILL, True),
        (
            "Obciążenia bez faktury z natury (opłaty, lokaty, podatki, ZUS, wynagrodzenia, "
            "przelewy własne)",
            bank["other"],
            NEUTRAL_FILL,
            False,
        ),
    ]
    row = 4
    for name, rows, fill, with_total in sections:
        rows = sorted(rows, key=lambda r: r["booking_date"])
        ws.cell(row=row, column=1, value=f"{name}: {len(rows)}").font = SECTION
        row += 1
        _header_row(ws, row, [c[0] for c in DEBIT_COLUMNS])
        row += 1
        for debit in rows:
            _debit_row(ws, row, debit, fill, missing=fill is MISSING_FILL)
            row += 1
        if rows and with_total:
            ws.cell(row=row, column=2, value="Razem").font = BOLD
            total = ws.cell(row=row, column=3, value=sum(r["amount"] for r in rows))
            total.font = BOLD
            total.number_format = MONEY
            row += 1
        row += 1
    _widths(ws, [c[1] for c in DEBIT_COLUMNS])
    ws.freeze_panes = "A4"


def _debit_row(ws, row: int, debit: dict, fill: PatternFill, missing: bool) -> None:
    invoice = debit["invoice"]
    description = (
        f"{debit['counterparty']} | {debit['description']}"
        if debit["counterparty"]
        else debit["description"]
    )
    values = [
        _date(debit["booking_date"]),
        _date(debit["operation_date"]),
        debit["amount"],
        debit["currency"],
        debit["kind"],
        description,
        invoice["number"] if invoice else ("BRAK FAKTURY" if missing else ""),
        _date(invoice["issue_date"]) if invoice else None,
        invoice["seller"] if invoice else None,
        invoice["gross"] if invoice else None,
        debit["method"],
        debit["note"],
    ]
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=col == 6)
    for col in (1, 2, 8):
        ws.cell(row=row, column=col).number_format = DATE
    for col in (3, 10):
        ws.cell(row=row, column=col).number_format = MONEY


def _invoices_sheet(ws, statuses: list[InvoiceStatus]) -> None:
    ws["A1"] = "Faktury miesiąca i sposób zapłaty"
    ws["A1"].font = TITLE
    ws["A2"] = (
        "Płatności znalezione na wyciągach z kolejnych miesięcy aktualizują to zestawienie. "
        "„Gotówka” – zapłacone z własnych środków (oznaczone ręcznie)."
    )
    _header_row(ws, 4, [c[0] for c in INVOICE_COLUMNS])
    order = list(STATUS_FILLS)
    rows = sorted(statuses, key=lambda s: (order.index(s.status), s.invoice.issue_date))
    for row, item in enumerate(rows, 5):
        invoice = item.invoice
        values = [
            invoice.issue_date,
            invoice.number,
            invoice.seller_name,
            invoice.seller_nip,
            invoice.gross,
            invoice.currency,
            item.status + (" (w kolejnym miesiącu)" if item.paid_later else ""),
            _date(item.paid_date),
            item.bank,
            item.statement_month,
            item.details,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value).fill = STATUS_FILLS[item.status]
        ws.cell(row=row, column=1).number_format = DATE
        ws.cell(row=row, column=5).number_format = MONEY
        ws.cell(row=row, column=8).number_format = DATE
    _widths(ws, [c[1] for c in INVOICE_COLUMNS])
    ws.freeze_panes = "A5"


def _header_row(ws, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _widths(ws, widths: list[int]) -> None:
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _date(value: str | None) -> date | None:
    return date.fromisoformat(value) if value else None
